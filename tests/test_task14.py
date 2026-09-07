"""Automated Test Suite for Task 14: Inspection Report Upload & Retrievable Cloud Links.

Verifies:
1. upload_inspection_report function in src/scan_service.py:
   - Uploads Word .docx report to Supabase Storage ('inspection-reports' bucket under reports/{scan_id}/inspection_report.docx).
   - Uploads PDF .pdf report (if provided) under reports/{scan_id}/inspection_report.pdf).
   - Retrieves public storage URLs (report_url, docx_url, pdf_url).
   - Updates 'report_url' in Supabase scans table via db.update_scan().
   - Returns ReportUploadResult container with report_url, pdf_url, upload_status, and scan update metadata.
   - Graceful offline fallback with mock storage and URLs.
   - Robust error handling for invalid/missing files and empty scan IDs.
   - Preservation of existing scan records and metadata upon update.
2. End-to-end integration:
   - Integrates with generate_inspection_report from src/report_generator.py.
   - Confirms storage content bytes match generated files.
   - Verifies database consistency via db.get_scan(scan_id).
3. Streamlit App Integration:
   - Headless AppTest verifying the automatic upload trigger upon report generation.
   - Verification of UI cloud storage link badges and retrievable URLs.
"""

from __future__ import annotations

import os
import shutil
import sys
import unittest
import uuid
import docx

# Ensure project root is in sys.path
PROJECT_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if PROJECT_ROOT not in sys.path:
    sys.path.insert(0, PROJECT_ROOT)

from src import db
from src.consolidation import ConsolidatedProductRecord
from src.report_generator import generate_inspection_report
from src.scan_service import (
    REPORT_STORAGE_BUCKET,
    REPORT_STORAGE_PREFIX,
    ReportUploadResult,
    upload_inspection_report,
)

TEST_TMP_DIR = os.path.join(PROJECT_ROOT, "tests", "temp_reports_task14")


class TestTask14ReportUpload(unittest.TestCase):
    """Unit and integration test cases for Task 14: Inspection Report Upload."""

    def setUp(self):
        """Set up temporary test directory and artifacts."""
        os.makedirs(TEST_TMP_DIR, exist_ok=True)
        self.test_scan_id = str(uuid.uuid4())

        # Create a real dummy .docx file using python-docx
        self.sample_docx_path = os.path.join(TEST_TMP_DIR, f"sample_report_{self.test_scan_id}.docx")
        doc = docx.Document()
        doc.add_heading("Legal Metrology Statutory Inspection Report", level=0)
        doc.add_paragraph(f"Scan ID: {self.test_scan_id}")
        doc.add_paragraph("Compliance Status: Fully Verified")
        doc.save(self.sample_docx_path)

        # Create a sample .pdf dummy file
        self.sample_pdf_path = os.path.join(TEST_TMP_DIR, f"sample_report_{self.test_scan_id}.pdf")
        with open(self.sample_pdf_path, "wb") as f_pdf:
            f_pdf.write(b"%PDF-1.4\n%Dummy PDF inspection report for Task 14 verification\n%%EOF\n")

    def tearDown(self):
        """Clean up temporary test artifacts."""
        if os.path.exists(TEST_TMP_DIR):
            shutil.rmtree(TEST_TMP_DIR, ignore_errors=True)

    def test_01_upload_inspection_report_docx_only(self):
        """Verify upload of Word .docx report only and database update."""
        scan_id = str(uuid.uuid4())
        res = upload_inspection_report(scan_id=scan_id, docx_path=self.sample_docx_path)

        # Validate return container
        self.assertIsInstance(res, ReportUploadResult)
        self.assertIsInstance(res, dict)
        self.assertEqual(res.status, "success")
        self.assertEqual(res.scan_id, scan_id)
        self.assertTrue(res.report_url)
        self.assertIn(f"reports/{scan_id}/inspection_report.docx", res.report_url)
        self.assertIsNone(res.pdf_url)
        self.assertEqual(res.storage_bucket, "inspection-reports")
        self.assertEqual(res.docx_storage_path, f"reports/{scan_id}/inspection_report.docx")

        # Verify database record updated via db.get_scan
        db_scan = db.get_scan(scan_id)
        self.assertIsNotNone(db_scan, "Scan record should exist in database.")
        self.assertEqual(db_scan["report_url"], res.report_url)

        # Verify storage content in mock store
        if db.is_offline_mode():
            self.assertIn("inspection-reports", db._mock_db.storage)
            self.assertIn(res.docx_storage_path, db._mock_db.storage["inspection-reports"])
            stored_bytes = db._mock_db.storage["inspection-reports"][res.docx_storage_path]
            with open(self.sample_docx_path, "rb") as f_orig:
                self.assertEqual(stored_bytes, f_orig.read())

    def test_02_upload_inspection_report_docx_and_pdf(self):
        """Verify simultaneous upload of both .docx and .pdf reports."""
        scan_id = str(uuid.uuid4())
        res = upload_inspection_report(
            scan_id=scan_id,
            docx_path=self.sample_docx_path,
            pdf_path=self.sample_pdf_path,
        )

        self.assertEqual(res.status, "success")
        self.assertIsNotNone(res.report_url)
        self.assertIsNotNone(res.pdf_url)
        self.assertIn(f"reports/{scan_id}/inspection_report.docx", res.report_url)
        self.assertIn(f"reports/{scan_id}/inspection_report.pdf", res.pdf_url)
        self.assertEqual(res.pdf_storage_path, f"reports/{scan_id}/inspection_report.pdf")

        # Verify database record reflects report_url
        db_scan = db.get_scan(scan_id)
        self.assertIsNotNone(db_scan)
        self.assertEqual(db_scan["report_url"], res.report_url)

        # Verify mock store has both files
        if db.is_offline_mode():
            self.assertIn(res.docx_storage_path, db._mock_db.storage["inspection-reports"])
            self.assertIn(res.pdf_storage_path, db._mock_db.storage["inspection-reports"])
            with open(self.sample_pdf_path, "rb") as f_orig_pdf:
                self.assertEqual(
                    db._mock_db.storage["inspection-reports"][res.pdf_storage_path],
                    f_orig_pdf.read(),
                )

    def test_03_existing_scan_record_update_and_preservation(self):
        """Verify updating existing scan preserves existing metadata and updates report_url."""
        scan_id = str(uuid.uuid4())

        # Pre-seed scan in database
        initial_scan = db.create_scan({
            "id": scan_id,
            "scan_id": scan_id,
            "status": "uploaded",
            "front_image_url": f"mock://product-images/scans/{scan_id}/front.png",
            "back_image_url": f"mock://product-images/scans/{scan_id}/back.png",
            "dimensions": {"summary": "Front: 1000x800 px"},
            "notes": "Pre-existing packaging scan",
        })
        self.assertIsNotNone(initial_scan)

        # Upload report
        res = upload_inspection_report(
            scan_id=scan_id,
            docx_path=self.sample_docx_path,
            pdf_path=self.sample_pdf_path,
        )

        # Query updated scan
        updated = db.get_scan(scan_id)
        self.assertIsNotNone(updated)
        self.assertEqual(updated["report_url"], res.report_url)
        # Ensure previous metadata was not overwritten
        self.assertEqual(updated["front_image_url"], f"mock://product-images/scans/{scan_id}/front.png")
        self.assertEqual(updated["back_image_url"], f"mock://product-images/scans/{scan_id}/back.png")
        self.assertEqual(updated["dimensions"]["summary"], "Front: 1000x800 px")
        self.assertIn("updated_at", updated)

    def test_04_error_handling_and_input_validations(self):
        """Verify proper validation and error reporting on invalid parameters."""
        # Missing or empty scan_id
        with self.assertRaises(ValueError):
            upload_inspection_report(scan_id="", docx_path=self.sample_docx_path)

        with self.assertRaises(ValueError):
            upload_inspection_report(scan_id=None, docx_path=self.sample_docx_path)

        # Missing or empty docx_path
        with self.assertRaises(ValueError):
            upload_inspection_report(scan_id=self.test_scan_id, docx_path="")

        with self.assertRaises(ValueError):
            upload_inspection_report(scan_id=self.test_scan_id, docx_path=None)

        # Non-existent docx file on disk
        with self.assertRaises(FileNotFoundError):
            upload_inspection_report(
                scan_id=self.test_scan_id,
                docx_path=os.path.join(TEST_TMP_DIR, "non_existent.docx"),
            )

        # Non-existent pdf_path should log warning and set pdf_url to None without crashing
        res = upload_inspection_report(
            scan_id=self.test_scan_id,
            docx_path=self.sample_docx_path,
            pdf_path=os.path.join(TEST_TMP_DIR, "non_existent.pdf"),
        )
        self.assertEqual(res.status, "success")
        self.assertIsNone(res.pdf_url)
        self.assertIsNotNone(res.report_url)

    def test_05_report_upload_result_dict_and_attribute_polymorphism(self):
        """Verify ReportUploadResult supports both dict key and attribute access."""
        res = upload_inspection_report(
            scan_id=self.test_scan_id,
            docx_path=self.sample_docx_path,
            pdf_path=self.sample_pdf_path,
        )

        # Attribute access
        self.assertEqual(res.report_url, res.docx_url)
        self.assertIn("mock://", res.report_url)
        self.assertIn("inspection_report.docx", res.docx_storage_path)
        self.assertIn("inspection_report.pdf", res.pdf_storage_path)
        self.assertEqual(res.status, "success")
        self.assertTrue(res.upload_status)

        # Dictionary key access
        self.assertEqual(res["report_url"], res.report_url)
        self.assertEqual(res["pdf_url"], res.pdf_url)
        self.assertEqual(res["scan_id"], self.test_scan_id)
        self.assertEqual(res["storage_bucket"], "inspection-reports")
        self.assertEqual(res["status"], "success")

        # Repr
        repr_str = repr(res)
        self.assertIn("ReportUploadResult", repr_str)
        self.assertIn(self.test_scan_id, repr_str)

    def test_06_end_to_end_generate_and_upload_report(self):
        """Verify complete pipeline from report_generator through upload_inspection_report."""
        scan_id = str(uuid.uuid4())
        record = ConsolidatedProductRecord(
            product_id=str(uuid.uuid4()),
            brand_name="Patanjali Cow Ghee",
            mrp="650.00",
            net_quantity="1 L",
            mfg_date="08/2026",
            expiry_date="08/2027",
            batch_number="PG-882",
            manufacturer_details="Patanjali Ayurved Ltd, Haridwar, Uttarakhand",
            consumer_care="customercare@patanjaliayurved.net",
            country_of_origin="India",
            unit_sale_price="Rs 0.65 / ml",
            completeness_pct=100.0,
            missing_declarations=[],
        )

        # 1. Generate inspection report (.docx and .pdf fallback)
        rep_res = generate_inspection_report(
            consolidated_record=record,
            output_dir=TEST_TMP_DIR,
            scan_id=scan_id,
        )
        self.assertTrue(os.path.isfile(rep_res.docx_path))

        # 2. Upload generated report
        upload_res = upload_inspection_report(
            scan_id=scan_id,
            docx_path=rep_res.docx_path,
            pdf_path=rep_res.pdf_path,
        )

        self.assertEqual(upload_res.status, "success")
        self.assertIn(f"reports/{scan_id}/inspection_report.docx", upload_res.report_url)

        if rep_res.pdf_path and os.path.exists(rep_res.pdf_path):
            self.assertIsNotNone(upload_res.pdf_url)
            self.assertIn(f"reports/{scan_id}/inspection_report.pdf", upload_res.pdf_url)

        # 3. Query DB to confirm persistence
        db_record = db.get_scan(scan_id)
        self.assertIsNotNone(db_record)
        self.assertEqual(db_record["report_url"], upload_res.report_url)


def test_streamlit_app_report_cloud_upload_and_ui_badge():
    """Verify Streamlit app report generation triggers upload and displays cloud link badge."""
    print("\n--- Testing Streamlit App Report Cloud Upload & Link Badge (AppTest) ---")
    try:
        from streamlit.testing.v1 import AppTest

        app_path = os.path.join(PROJECT_ROOT, "app.py")
        at = AppTest.from_file(app_path, default_timeout=35)
        at.run()

        test_scan_id = str(uuid.uuid4())
        at.session_state["active_scan_id"] = test_scan_id
        at.session_state["current_scan"] = {
            "scan_id": test_scan_id,
            "timestamp": "2026-09-07T13:30:00Z",
            "status": "processed",
            "storage_status": "Offline Mock",
            "is_offline": True,
            "front_image_url": f"mock://product-images/scans/{test_scan_id}/front.png",
            "back_image_url": f"mock://product-images/scans/{test_scan_id}/back.png",
            "dimensions": {"summary": "Front: 600x400 px | Back: 600x400 px"},
        }

        # Inject consolidated record
        sample_rec = ConsolidatedProductRecord(
            product_id=str(uuid.uuid4()),
            brand_name="Tata Salt Super-Refined",
            mrp="28.00",
            net_quantity="1 kg",
            mfg_date="06/2026",
            expiry_date="06/2028",
            batch_number="TS-2026",
            manufacturer_details="Tata Chemicals Ltd, Gujarat",
            consumer_care="feedback@tatachemicals.com",
            country_of_origin="India",
            unit_sale_price="Rs 0.028/g",
            completeness_pct=100.0,
            missing_declarations=[],
        )
        at.session_state["consolidated_record"] = sample_rec

        # Re-run app to render Consolidated Product Master Record and report button
        at.run()
        assert not at.exception, f"AppTest exception on render: {at.exception}"

        # Locate '📄 Generate Inspection Report' button
        gen_buttons = [b for b in at.button if "Generate Inspection Report" in b.label]
        assert gen_buttons, f"Button '📄 Generate Inspection Report' not found: {[b.label for b in at.button]}"
        print("✓ Streamlit '📄 Generate Inspection Report' button located")

        # Click Generate Inspection Report button
        gen_buttons[0].click().run()
        assert not at.exception, f"AppTest exception after button click: {at.exception}"

        # 1. Verify automatic upload occurred and is stored in session state
        assert "report_upload_result" in at.session_state, "report_upload_result not found in session_state"
        upload_result = at.session_state["report_upload_result"]
        assert upload_result is not None, "session_state['report_upload_result'] is None"
        assert upload_result.report_url, "upload_result.report_url is empty"
        print(f"✓ Automatic upload confirmed: {upload_result.report_url}")

        # 2. Verify database scan record was updated
        db_scan = db.get_scan(test_scan_id)
        assert db_scan is not None, f"Scan {test_scan_id} not found in database"
        assert db_scan.get("report_url") == upload_result.report_url, (
            f"Database report_url mismatch: expected {upload_result.report_url}, got {db_scan.get('report_url')}"
        )
        print(f"✓ Database scan record updated with report_url: {db_scan.get('report_url')}")

        # 3. Verify UI cloud link badge and elements are rendered
        markdown_texts = [m.value for m in at.markdown]
        badge_rendered = any("Cloud Storage" in t or "Retrievable Inspection Report Links" in t for t in markdown_texts)
        assert badge_rendered, f"Cloud storage badge not found in markdown: {markdown_texts}"
        print("✓ Streamlit UI cloud storage link badge verified")

        info_texts = [i.value for i in at.info]
        assert any("Cloud Word Report" in t and upload_result.report_url in t for t in info_texts), (
            f"Cloud Word Report info element missing from: {info_texts}"
        )
        print("✓ Streamlit UI retrievable cloud Word report link element verified")

        if upload_result.pdf_url:
            assert any("Cloud PDF Report" in t and upload_result.pdf_url in t for t in info_texts), (
                f"Cloud PDF Report info element missing from: {info_texts}"
            )
            print("✓ Streamlit UI retrievable cloud PDF report link element verified")

    except ImportError:
        print("ℹ️ streamlit.testing.v1 not available; skipping AppTest")


if __name__ == "__main__":
    print("============================================================")
    print("STARTING TASK 14 VERIFICATION SUITE")
    print("============================================================")
    suite = unittest.TestLoader().loadTestsFromTestCase(TestTask14ReportUpload)
    runner = unittest.TextTestRunner(verbosity=2)
    result = runner.run(suite)
    assert result.wasSuccessful(), f"Task 14 unit tests failed: {result.failures or result.errors}"

    # Run Streamlit App headless integration test
    test_streamlit_app_report_cloud_upload_and_ui_badge()
    print("\n============================================================")
    print("ALL TASK 14 TESTS PASSED SUCCESSFULLY (0 FAILURES)")
    print("============================================================")
